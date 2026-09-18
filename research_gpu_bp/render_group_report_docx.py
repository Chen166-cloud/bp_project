"""Create the 2026-09-19 BP/FFBP group-meeting report as an editable DOCX."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "output" / "docx" / "9.19组会.docx"
FIGURE_1 = ROOT / "fast_bp" / "results" / "ffbp_final_preview.png"
FIGURE_2 = ROOT / "fast_bp" / "results" / "gpu_baseline" / "bp_ffbp_cpu_gpu_fourway.png"
FIGURE_3 = ROOT / "fast_bp" / "results" / "gpu_baseline" / "bp_ffbp_cpu_gpu_timing.png"

BLUE = "1F4E78"
PALE_BLUE = "F2F6F9"
GRID = "000000"
TEXT = RGBColor(0, 0, 0)
GRAY = RGBColor(89, 89, 89)


def set_east_asia_font(run, name: str, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.color.rgb = TEXT
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_style_font(style, name: str, size: float, bold: bool = False) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = TEXT
    rpr = style.element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:eastAsia"), name)
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)


def remove_paragraph_borders(paragraph_or_style) -> None:
    element = paragraph_or_style._p if hasattr(paragraph_or_style, "_p") else paragraph_or_style.element
    p_pr = element.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 100, start: int = 100, bottom: int = 100, end: int = 100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, margin_value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "5")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), GRID)


def set_cell_text(cell, text: str, *, header: bool = False, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.12
    paragraph.clear()
    match = re.fullmatch(r"(.+?)\^(-?\d+)", text)
    if match:
        base = paragraph.add_run(match.group(1))
        exponent = paragraph.add_run(match.group(2))
        exponent.font.superscript = True
        for run in (base, exponent):
            set_east_asia_font(run, "微软雅黑", 8.2, header)
    else:
        run = paragraph.add_run(text)
        set_east_asia_font(run, "微软雅黑", 8.2 if not header else 8.7, header)
    if header:
        for run in paragraph.runs:
            run.font.color.rgb = TEXT
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)


def format_table(table, widths: list[float], numeric_columns: set[int] | None = None) -> None:
    numeric_columns = numeric_columns or set()
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        prevent_row_split(row)
        for column_index, cell in enumerate(row.cells):
            cell.width = Cm(widths[column_index])
            if row_index == 0:
                shade_cell(cell, "FFFFFF")
                align = WD_ALIGN_PARAGRAPH.CENTER
            else:
                shade_cell(cell, "FFFFFF")
                align = WD_ALIGN_PARAGRAPH.CENTER if column_index in numeric_columns else WD_ALIGN_PARAGRAPH.LEFT
            existing_text = cell.text
            set_cell_text(cell, existing_text, header=row_index == 0, align=align)


def add_heading(document: Document, text: str, level: int) -> None:
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True


def add_body(document: Document, text: str, *, first_line: bool = True) -> None:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.first_line_indent = Cm(0.74) if first_line else Cm(0)
    paragraph.paragraph_format.keep_together = True
    parts = re.split(r"(\d+(?:\.\d+)?×10\^-?\d+)", text)
    for part in parts:
        scientific = re.fullmatch(r"(.+?×10)\^(-?\d+)", part)
        if scientific:
            base = paragraph.add_run(scientific.group(1))
            exponent = paragraph.add_run(scientific.group(2))
            exponent.font.superscript = True
            set_east_asia_font(base, "宋体", 10.5)
            set_east_asia_font(exponent, "宋体", 10.5)
        elif part:
            run = paragraph.add_run(part)
            set_east_asia_font(run, "宋体", 10.5)


def add_figure(document: Document, path: Path, width_cm: float, caption: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(path), width=Cm(width_cm))

    caption_paragraph = document.add_paragraph(style="Caption")
    caption_paragraph.paragraph_format.keep_together = True
    caption_paragraph.paragraph_format.space_after = Pt(7)
    label, remainder = caption.split(" ", 1)
    label_run = caption_paragraph.add_run(label + " ")
    set_east_asia_font(label_run, "微软雅黑", 9, True)
    body_run = caption_paragraph.add_run(remainder)
    set_east_asia_font(body_run, "微软雅黑", 9)


def add_reference(document: Document, number: int, text: str, url: str | None = None) -> None:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.left_indent = Cm(0.7)
    paragraph.paragraph_format.first_line_indent = Cm(-0.7)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(f"{number}. {text}")
    set_east_asia_font(run, "微软雅黑", 9)
    if url:
        add_hyperlink(paragraph, url, url)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    new_run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(color)
    run_properties.append(underline)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:eastAsia"), "微软雅黑")
    fonts.set(qn("w:ascii"), "微软雅黑")
    fonts.set(qn("w:hAnsi"), "微软雅黑")
    run_properties.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    run_properties.append(size)
    new_run.append(run_properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_page_number(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    prefix = paragraph.add_run("第 ")
    set_east_asia_font(prefix, "微软雅黑", 8)
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    field_run = paragraph.add_run()
    field_run._r.append(field_begin)
    field_run._r.append(instruction)
    field_run._r.append(field_end)
    suffix = paragraph.add_run(" 页")
    set_east_asia_font(suffix, "微软雅黑", 8)


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    set_style_font(normal, "宋体", 10.5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    title = styles["Title"]
    set_style_font(title, "微软雅黑", 21, True)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(16)
    remove_paragraph_borders(title)

    heading_1 = styles["Heading 1"]
    set_style_font(heading_1, "微软雅黑", 15, True)
    heading_1.paragraph_format.space_before = Pt(12)
    heading_1.paragraph_format.space_after = Pt(6)
    heading_1.paragraph_format.keep_with_next = True

    heading_2 = styles["Heading 2"]
    set_style_font(heading_2, "微软雅黑", 12.5, True)
    heading_2.paragraph_format.space_before = Pt(9)
    heading_2.paragraph_format.space_after = Pt(5)
    heading_2.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    set_style_font(caption, "微软雅黑", 9)
    caption.font.color.rgb = GRAY
    caption.paragraph_format.line_spacing = 1.15
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


def build_document() -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    section.footer_distance = Cm(0.7)
    configure_styles(document)
    add_page_number(section)

    document.core_properties.title = "2026年9月19日 BP与FFBP CPU GPU成像组会汇报"
    document.core_properties.subject = "CPU与GPU直接BP和FFBP成像基线比较"
    document.core_properties.author = "SAR imaging project"

    title = document.add_paragraph(style="Title")
    remove_paragraph_borders(title)
    title_run = title.add_run("2026年9月19日 BP与FFBP成像组会汇报")
    set_east_asia_font(title_run, "微软雅黑", 21, True)

    add_body(
        document,
        "本次汇报比较直接BP和FFBP在CPU与GPU上的成像结果、复数误差和驻留计算时间。固定混合精度GPU保持了同算法CPU结果，当前主要质量差异来自FFBP的分解与重采样近似。",
        first_line=False,
    )

    add_heading(document, "算法来源", 1)
    source_table = document.add_table(rows=3, cols=2)
    source_rows = [
        ["算法", "出处与本实验实现"],
        [
            "BP",
            "SAR层析反投影基础见[1]；本实验所用直接时域BP形式通常追溯至McCorkle[2]。逐像素计算双程距离，对距离压缩回波插值、相位补偿并沿孔径相干累加。",
        ],
        [
            "FFBP",
            "采用Ulander、Hellsten和Stenström[3]的快速因子化反投影思想：将孔径递归分组，在局部极坐标网格形成子图，经重采样后逐层合并，以插值近似换取较低计算量。",
        ],
    ]
    for row, values in zip(source_table.rows, source_rows):
        for cell, value in zip(row.cells, values):
            cell.text = value
    format_table(source_table, [2.3, 14.9], numeric_columns={0})

    add_heading(document, "2 全场成像与四种基线", 1)
    add_heading(document, "2.1 自聚焦后FFBP全场景地面图", 2)
    add_figure(
        document,
        FIGURE_1,
        16.3,
        "图1 2018-08-14_18-54-08.dat的FFBP/PGA自聚焦地面成像结果。横向显示范围为450至2650 m，航向范围为-800至800 m，幅度窗为-65至-10 dB。该图对应全场景FFBP结果，不作为四种CPU/GPU基线的同网格定量对照。",
    )

    document.add_page_break()
    add_heading(document, "2.2 CPU与GPU四种基线", 2)
    add_figure(
        document,
        FIGURE_2,
        15.5,
        "图2 同一实测ROI的四种成像结果。上行为直接BP，下行为FFBP；左列为CPU，右列为固定混合精度GPU。四幅图以CPU直接BP峰值为共同参考，显示范围均为-50至0 dB。图中为极坐标复图像的幅度，不是地面投影图。",
    )

    result_table = document.add_table(rows=5, cols=6)
    result_rows = [
        ["实验", "驻留计算时间", "数值参考", "相对复误差", "SER / dB", "峰值像素"],
        ["CPU直接BP", "10,200.507 ms", "高精度基准", "-", "-", "[222, 201]"],
        ["GPU直接BP 固定混合", "611.359 ms", "CPU直接BP", "1.879×10^-6", "114.52", "[222, 201]"],
        ["CPU FFBP", "538.205 ms", "CPU直接BP", "3.2858%", "29.67", "[222, 201]"],
        ["GPU FFBP 固定混合", "50.085 ms", "CPU FFBP", "2.040×10^-7", "133.81", "[222, 201]"],
    ]
    for row, values in zip(result_table.rows, result_rows):
        for cell, value in zip(row.cells, values):
            cell.text = value
    format_table(result_table, [3.2, 2.7, 2.7, 3.0, 2.2, 2.3], numeric_columns={1, 3, 4, 5})

    add_body(
        document,
        "GPU相对CPU的同算法加速比分别为16.68×（直接BP）和10.75×（FFBP）。同一设备上，FFBP相对直接BP的加速比为18.95×（CPU）和12.21×（GPU）。GPU FFBP相对CPU直接BP达到203.66×，但其相对CPU直接BP的总SER为29.67 dB，因此该倍率不属于等质量加速。",
    )
    add_body(
        document,
        "GPU直接BP与GPU FFBP对各自CPU参考的误差均低于2×10^-6，说明CUDA实现保持了既定成像算子。CPU FFBP相对CPU直接BP的3.29%复误差反映因子分解和局部网格重采样的算法近似。",
    )

    document.add_page_break()
    add_figure(
        document,
        FIGURE_3,
        16.0,
        "图3 四种基线的核心成像时间，纵轴为对数尺度。GPU时间不含文件读取、H2D和D2H；该口径对应多区域、多帧或参数扫描时复用显存中数据的计算场景。",
    )
    add_body(
        document,
        "固定混合精度GPU复现了同算法CPU结果：直接BP和FFBP的驻留计算时间分别由10.201 s、538.205 ms降至611.359 ms、50.085 ms。FFBP的主要误差来自分解与重采样近似，而非GPU移植。",
    )
    add_body(
        document,
        "若计入单次H2D和D2H，GPU直接BP为1628.993 ms，仍比CPU快6.26×；GPU FFBP为1067.767 ms，慢于CPU FFBP的538.205 ms。纯FP32直接BP为38.050 ms / 48.67 dB，FFBP为45.877 ms / 38.78 dB。",
    )
    add_body(
        document,
        "本轮评价对象为计算时间和复数一致性。当前实测场景没有已标定点目标，空间分辨率、IRW、PSLR与ISLR留待点目标实验测量。",
    )

    add_heading(document, "3 结论", 1)
    conclusions = [
        "两种GPU实现均通过同算法CPU复图核对，峰值位置一致，固定混合精度可作为后续优化的正确性基线。",
        "数据常驻条件下，GPU对直接BP和FFBP分别加速16.68×和10.75×；四种配置中GPU FFBP用时最短。",
        "当前质量差异主要由FFBP的分解与重采样决定。后续比较应在相同输出网格和相同误差约束下进行。",
    ]
    for conclusion in conclusions:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.keep_together = True
        run = paragraph.add_run(conclusion)
        set_east_asia_font(run, "宋体", 10.5)

    add_heading(document, "4 接下来的工作", 1)
    add_body(
        document,
        "下一步考虑用Nsight Systems分析按层批处理FFBP的耗时，查明时间主要花在叶节点BP、重采样还是核函数调度上。在相同输出网格下调整叶节点大小、内部采样密度和各层计算精度，比较SER不低于45 dB时的计算速度。成像质量还需结合亮目标相位误差及点目标IRW、PSLR、ISLR判断。实验先使用现有实测数据和点目标仿真，再补充GOTCHA数据及其他型号GPU的对照实验。",
    )

    document.add_page_break()
    add_heading(document, "参考文献", 1)
    add_reference(
        document,
        1,
        "Munson, D. C., O'Brien, J. D., Jenkins, W. K. A Tomographic Formulation of Spotlight-Mode Synthetic Aperture Radar. Proceedings of the IEEE, 1983, 71(8): 917-925.",
    )
    add_reference(
        document,
        2,
        "McCorkle, J. W. Focusing of Synthetic Aperture Ultra Wideband Data. IEEE International Conference on Systems Engineering, 1991: 1-5.",
    )
    add_reference(
        document,
        3,
        "Ulander, L. M. H., Hellsten, H., Stenström, G. Synthetic-Aperture Radar Processing Using Fast Factorized Back-Projection. IEEE Transactions on Aerospace and Electronic Systems, 2003, 39(3): 760-776.",
    )
    add_reference(
        document,
        4,
        "Fasih, A., Hartley, T. GPU-Accelerated Synthetic Aperture Radar Backprojection in CUDA. IEEE Radar Conference, 2010.",
    )
    add_reference(document, 5, "NVIDIA. MatX SAR Back Projection API. ", "https://nvidia.github.io/MatX/api/signalimage/radar/sar_bp.html")
    add_reference(document, 6, "Benson, T. GOTCHA-Back: CUDA Backprojection Implementation and SER Comparison. ", "https://github.com/tbensonatl/gotcha-back")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
